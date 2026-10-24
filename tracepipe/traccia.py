"""One record shape for the three corpora of TRACE (``traccia``: trace).

TRACE-J (judicial), TRACE-N (laws and policies) and TRACE-M (media) keep their working files in three shapes: the
citations of the Zenodo release and of ``corpus_rows``, the rows of the S07-N name search on CCLW, the articles of the
media index. This module turns each of them into the record of ``schema/trace-record.schema.json``, which is what the
public files and the public API give out, and checks it against the schema. The working files are never rewritten.

What the record adds to the working files:

- the identifiers other databases join on: the Climate Policy Radar ids of the citing document and of its case or
  family (``Sabin.document.N.M``, ``Sabin.family.N.0``), the bundles, the ISO 3166 codes of the geography, and the
  CLIMPACT ids of the cited case and decision (``identifiers/``);
- an ``anchor`` that identifies an occurrence without ``char_start``, which moves whenever the text is rebuilt;
- the access tier and the licence on every record, so that the rule «no media text» is a property of the data and not
  of each page that shows it;
- text with no control characters (one NUL byte in the 0.1 citations file stops the CSV readers).

``trace traccia converti`` writes the records, ``trace traccia valida`` checks a file of records.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from collections.abc import Iterable, Iterator, Mapping
from pathlib import Path
from typing import Any, ClassVar

from .righe import normalizza

SCHEMA_VERSIONE = '0.1.0'
SCHEMA = Path(__file__).resolve().parent.parent / 'schema' / 'trace-record.schema.json'

CORPORA: dict[str, dict[str, str]] = {
    'J': {'fonte': 'climate-litigation-database', 'nome': 'TRACE-J'},
    'N': {'fonte': 'climate-change-laws-of-the-world', 'nome': 'TRACE-N'},
    'M': {'fonte': 'nexis', 'nome': 'TRACE-M'},
}
"""The corpus letters. Renaming a corpus (its letter or its public name) is a change of this table and of the
``corpus`` enum of the schema, with the old letter kept as an alias of the codes."""

LICENZA_J = 'CC BY 4.0 (TRACE data); citing texts © their authors, from the Climate Litigation Database'
ATTRIBUZIONE_J = 'Data provided by the Sabin Center for Climate Change Law. Platform operated by Climate Policy Radar.'
LICENZA_N = 'CC BY 4.0 (TRACE data); citing texts © their authors, from Climate Change Laws of the World'
ATTRIBUZIONE_N = 'Climate Change Laws of the World, Grantham Research Institute and Climate Policy Radar.'
LICENZA_M = 'metadata only; texts under the Nexis Uni licence of the University of Bari, not redistributed'

_CONTROLLO = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f]')
_DATA = re.compile(r'^\d{4}(-\d{2}(-\d{2})?)?$')
_CODICE = re.compile(r'^([A-Z]{2,5})-(\d{1,6})(?:-([JNM]))?$')
_STATO_J = {'checked by the team': 'checked', 'proposed by the machine, not yet checked': 'proposed'}

Riga = Mapping[str, Any]


class Pulizia(Counter):
    """How many control characters were removed, by field: reported, never silent."""


def testo(valore: Any, campo: str = '', pulizia: Pulizia | None = None) -> str | None:
    """The text without control characters; ``None`` for an empty value."""
    if valore is None:
        return None
    s = str(valore)
    pulito = _CONTROLLO.sub('', s)
    if pulizia is not None and pulito != s:
        pulizia[campo] += 1
    return pulito.strip() or None


def lista(valore: Any) -> list[str]:
    """A list from a real list or from the «; »-joined form of the CSV files."""
    if valore is None or valore == '':
        return []
    if isinstance(valore, list):
        return [str(x).strip() for x in valore if str(x).strip()]
    return [x.strip() for x in re.split(r'\s*;\s*', str(valore)) if x.strip()]


def intero(valore: Any) -> int | None:
    try:
        return int(str(valore).strip())
    except (TypeError, ValueError):
        return None


def data(valore: Any) -> str | None:
    s = str(valore or '').strip()[:10]
    return s if _DATA.match(s) else None


def ancora(
    documento: str, pagina: int | None, prefisso: str, frase: str | None, passo: str | None = None
) -> str | None:
    """``<document>|<page>|<cited prefix>|<16 hex of the sentence>|<8 hex of the passage>``, or ``None`` without a page
    or any text. The SHA-256 is taken on the normalised text, so the anchor survives a rebuild of the text (OCR,
    Markdown, headers) that moves ``char_start``. The cited prefix is in it because one sentence can cite two cases;
    the passage is in it because two footnotes on one page can share the same extracted sentence."""
    if not documento or not pagina or not (frase or passo):
        return None

    def h(t: str | None, n: int) -> str:
        return hashlib.sha256(normalizza(t or '').encode('utf-8')).hexdigest()[:n]

    return f'{documento}|{pagina}|{prefisso}|{h(frase, 16)}|{h(passo, 8)}'


def codice_con_lettera(codice: str, corpus: str) -> tuple[str, list[str]]:
    """The code with its corpus letter, and the earlier form as alias (URG-0146 -> URG-0146-J, [URG-0146])."""
    m = _CODICE.match(str(codice).strip().upper())
    if not m:
        raise ValueError(f'not a citation code: {codice!r}')
    prefisso, numero, lettera = m.groups()
    if lettera and lettera != corpus:
        raise ValueError(f'{codice}: letter {lettera}, corpus {corpus}')
    pieno = f'{prefisso}-{numero}-{corpus}'
    return pieno, ([] if lettera else [f'{prefisso}-{numero}'])


def iso(valore: Any) -> list[str]:
    return [x.upper() for x in lista(valore)]


# ── what the records are joined with ───────────────────────────────────────────────────────────────────────────────


class Identificativi:
    """The CLIMPACT ids of the cited cases and decisions (``identifiers/ids.csv``, ``citations_to_cases.csv``)."""

    def __init__(self, cartella: Path) -> None:
        self.caso_da_prefisso: dict[str, str] = {}
        with open(cartella / 'ids.csv', encoding='utf-8', newline='') as f:
            for r in csv.DictReader(f):
                if r['type'] == 'case' and r.get('citation_prefix') and r.get('status') == 'active':
                    self.caso_da_prefisso[r['citation_prefix']] = r['climpact_id']
        self.per_codice: dict[str, dict[str, str]] = {}
        p = cartella / 'citations_to_cases.csv'
        if p.exists():
            with open(p, encoding='utf-8', newline='') as f:
                for r in csv.DictReader(f):
                    self.per_codice[r['citation_code']] = r

    def caso(self, prefisso: str) -> str | None:
        return self.caso_da_prefisso.get(prefisso)


class MetadatiFonte:
    """The export of the Climate Litigation Database, by document slug: the Climate Policy Radar ids, bundles,
    languages and ISO codes that the 0.1 citations file leaves out.

    Each source is the ``metadata.csv`` of the corpus or a whole-database export as downloaded
    (``CCC-whole_data_dump-<date>.zip``, an .xlsx inside, read with ``openpyxl``). With several sources the later
    ones win: give them oldest first, so that a document added after the first export is still found."""

    COLONNE: ClassVar[dict[str, str]] = {
        'document_import_id': 'Internal Document ID',
        'family_id': 'Case ID',
        'family_import_id': 'Internal Case ID',
        'family_title': 'Case Name',
        'bundle_ids': 'Bundle ID(s)',
        'languages': 'Language(s)',
        'geography_iso': 'Geography ISOs',
    }

    def __init__(self, *percorsi: Path) -> None:
        self.per_documento: dict[str, dict[str, str]] = {}
        for p in percorsi:
            for r in self._righe(Path(p)):
                if r.get('Document ID'):
                    self.per_documento[r['Document ID']] = {k: str(r.get(c) or '') for k, c in self.COLONNE.items()}

    @staticmethod
    def _righe(p: Path) -> Iterator[dict[str, Any]]:
        if p.suffix == '.csv':
            csv.field_size_limit(sys.maxsize)
            with open(p, encoding='utf-8', newline='') as f:
                yield from csv.DictReader(f)
            return
        import io
        import zipfile

        from openpyxl import load_workbook

        with zipfile.ZipFile(p) as z:
            nome = next(n for n in z.namelist() if n.endswith('.xlsx'))
            libro = load_workbook(io.BytesIO(z.read(nome)), read_only=True)
        righe = libro.worksheets[0].iter_rows(values_only=True)
        intestazione = [str(x) for x in next(righe)]
        for valori in righe:
            yield dict(zip(intestazione, valori, strict=False))

    def __call__(self, documento: str) -> dict[str, str]:
        return self.per_documento.get(documento, {})


# ── the three converters ────────────────────────────────────────────────────────────────────────────────────────────


def da_j(r: Riga, ids: Identificativi, meta: MetadatiFonte, pulizia: Pulizia | None = None) -> dict[str, Any]:
    """A citation of the TRACE 0.1 dataset (``trace-0.1-citations.jsonl``) as a record of corpus J."""
    codice, alias = codice_con_lettera(r['citation_code'], 'J')
    prefisso = codice.split('-')[0]
    legame = ids.per_codice.get(str(r['citation_code']), {})
    m = meta(r['document_id'])
    pagina = intero(r.get('page'))
    frase = testo(r.get('sentence'), 'sentence', pulizia)
    passo = testo(r.get('passage'), 'passage', pulizia)
    stato = _STATO_J.get(str(r.get('review_status', '')), 'proposed')
    return {
        'schema_version': SCHEMA_VERSIONE,
        'corpus': 'J',
        'kind': 'occurrence',
        'ids': {
            'trace_id': None,
            'code': codice,
            'aliases': alias,
            'anchor': ancora(r['document_id'], pagina, prefisso, frase, passo),
        },
        'cited': {
            'case_id': legame.get('case_id') or ids.caso(prefisso),
            'decision_id': legame.get('decision_id') or None,
            'decision_basis': legame.get('decision_basis') or None,
            'code_prefix': prefisso,
            'label': testo(r.get('cited_judgment')) or prefisso,
        },
        'citing': {
            'source': CORPORA['J']['fonte'],
            'document_id': r['document_id'],
            'document_import_id': m.get('document_import_id') or None,
            'family_id': m.get('family_id') or None,
            'family_import_id': m.get('family_import_id') or None,
            'bundle_ids': lista(m.get('bundle_ids')),
            'title': testo(r.get('document_name')),
            'family_title': testo(m.get('family_title')),
            'document_type': testo(r.get('document_type')),
            'date': data(r.get('document_date')) or data(r.get('document_year')),
            'languages': lista(m.get('languages')),
            'geography_iso': iso(m.get('geography_iso')),
            'geography_label': testo(r.get('country')),
            'court': testo(r.get('court')),
            'outlet': None,
            'url': r.get('document_url') or None,
        },
        'location': {
            'page': pagina,
            'char_start': None,
            'text_version': r.get('corpus_version') or None,
            'where_in_text': r.get('where_in_text') or None,
            'form_found': None,
        },
        'evidence': {'sentence': frase, 'passage': passo},
        'classification': {
            'citation_function': r.get('citation_function') or None,
            'who_speaks': r.get('who_speaks') or None,
            'reference_kind': r.get('reference_kind') or None,
            'principles': [p for p in (testo(x) for x in lista(r.get('principles'))) if p],
            'cited_paragraphs': lista(r.get('cited_paragraphs')),
            'surrounding_pages': [n for n in (intero(x) for x in lista(r.get('surrounding_pages'))) if n],
            'level': None,
            'subjects': [],
            'relevance': None,
        },
        'provenance': {
            'found_by': r.get('found_by') or None,
            'corpus_version': r.get('corpus_version') or None,
            'proposal_rule': r.get('proposal_rule') or None,
            'five_readings': {'yes': True, 'no': False}.get(str(r.get('five_readings', '')).lower()),
            'link_rule': None,
            'query_fingerprint': None,
            'batch': None,
        },
        'review': {'status': stato},
        'access': {'tier': 2, 'licence': LICENZA_J, 'attribution': ATTRIBUZIONE_J},
    }


def da_n(r: Riga, ids: Identificativi, esecuzione: str, pulizia: Pulizia | None = None) -> dict[str, Any]:
    """A row of the S07-N name search on CCLW (``runs/S07-N-v3/<date>/righe_<case>.jsonl``) as a record of corpus N."""
    codice, alias = codice_con_lettera(r['codice'], 'N')
    prefisso = codice.split('-')[0]
    pagina = intero(r.get('pagina'))
    frase = testo(r.get('citazione_testuale'), 'citazione_testuale', pulizia)
    return {
        'schema_version': SCHEMA_VERSIONE,
        'corpus': 'N',
        'kind': 'occurrence',
        'ids': {
            'trace_id': None,
            'code': codice,
            'aliases': alias,
            'anchor': ancora(r['document_id'], pagina, prefisso, frase),
        },
        'cited': {
            'case_id': ids.caso(prefisso),
            'decision_id': None,
            'decision_basis': None,
            'code_prefix': prefisso,
            'label': testo(r.get('caso')) or prefisso,
        },
        'citing': {
            'source': CORPORA['N']['fonte'],
            'document_id': r['document_id'],
            'document_import_id': None,
            'family_id': r.get('family_id') or None,
            'family_import_id': None,
            'bundle_ids': [],
            'title': testo(r.get('document_title')),
            'family_title': testo(r.get('family_title')),
            'document_type': testo(r.get('tipo_atto')),
            'date': data(r.get('data')),
            'languages': [],
            'geography_iso': iso(r.get('geografie_iso')),
            'geography_label': testo(r.get('geografie')),
            'court': None,
            'outlet': None,
            'url': r.get('url_documento') or None,
        },
        'location': {
            'page': pagina,
            'char_start': intero(r.get('char_start')),
            'text_version': None,
            'where_in_text': None,
            'form_found': testo(r.get('forma_trovata')),
        },
        'evidence': {'sentence': frase, 'passage': None},
        'classification': None,
        'provenance': {
            'found_by': esecuzione,
            'corpus_version': None,
            'proposal_rule': None,
            'five_readings': None,
            'link_rule': testo(r.get('legame')),
            'query_fingerprint': None,
            'batch': None,
        },
        'review': {'status': 'proposed'},
        'access': {'tier': 2, 'licence': LICENZA_N, 'attribution': ATTRIBUZIONE_N},
    }


def da_m(a: Riga, ids: Identificativi, impronta: str | None = None) -> dict[str, Any]:
    """An article of the media index (``<PREFIX>-M/indice.json``) as a record of corpus M: metadata only, no text."""
    codice, alias = codice_con_lettera(a['codice'], 'M')
    prefisso = codice.split('-')[0]
    profilo = a.get('profilo_testata') or {}
    qid = profilo.get('wikidata_qid') or None
    soggetti = [{'name': str(s.get('nome')), 'weight': s.get('pct')} for s in (a.get('subject') or []) if s.get('nome')]
    pertinenza = a.get('pertinenza')
    if isinstance(pertinenza, Mapping):
        pertinenza = pertinenza.get('classe') or pertinenza.get('esito') or json.dumps(pertinenza, ensure_ascii=False)
    return {
        'schema_version': SCHEMA_VERSIONE,
        'corpus': 'M',
        'kind': 'item',
        'ids': {'trace_id': None, 'code': codice, 'aliases': alias, 'anchor': None},
        'cited': {
            'case_id': ids.caso(prefisso),
            'decision_id': None,
            'decision_basis': None,
            'code_prefix': prefisso,
            'label': testo(a.get('caso')) or prefisso,
        },
        'citing': {
            'source': CORPORA['M']['fonte'],
            'document_id': a.get('urn') or codice,
            'document_import_id': None,
            'family_id': None,
            'family_import_id': None,
            'bundle_ids': [],
            'title': testo(a.get('titolo')),
            'family_title': None,
            'document_type': testo(a.get('tipo_pubblicazione')),
            'date': data(a.get('data')),
            'languages': lista(a.get('lingua'))[:1],
            'geography_iso': [],
            'geography_label': testo(a.get('paese_testata')),
            'court': None,
            'outlet': {
                'name': testo(a.get('testata')) or '',
                'country': testo(profilo.get('paese') or a.get('paese_testata')),
                'scope': profilo.get('ambito') or None,
                'format': profilo.get('formato') or None,
                'publisher': profilo.get('editore') or None,
                'wikidata': qid if qid and re.match(r'^Q\d+$', qid) else None,
            },
            'url': None,
        },
        'location': None,
        'evidence': None,
        'classification': {
            'citation_function': None,
            'who_speaks': None,
            'reference_kind': None,
            'principles': [],
            'cited_paragraphs': [],
            'surrounding_pages': [],
            'level': None,
            'subjects': soggetti,
            'relevance': str(pertinenza) if pertinenza not in (None, '') else None,
        },
        'provenance': {
            'found_by': None,
            'corpus_version': None,
            'proposal_rule': None,
            'five_readings': None,
            'link_rule': None,
            'query_fingerprint': impronta,
            'batch': a.get('blocco') or None,
        },
        'review': {'status': 'proposed'},
        'access': {'tier': 3, 'licence': LICENZA_M, 'attribution': None},
    }


# ── the check ───────────────────────────────────────────────────────────────────────────────────────────────────────


def validatore(schema: Path = SCHEMA) -> Any:
    """A ``jsonschema`` validator of the record (``pip install -e '.[schema]'``)."""
    try:
        import jsonschema
    except ImportError as e:  # pragma: no cover - the dev extra installs it
        raise SystemExit("the check needs jsonschema: pip install -e '.[schema]'") from e
    s = json.loads(schema.read_text(encoding='utf-8'))
    return jsonschema.Draft202012Validator(s, format_checker=jsonschema.FormatChecker())


def errori(record: Mapping[str, Any], v: Any) -> list[str]:
    return [f'{"/".join(map(str, e.absolute_path)) or "(record)"}: {e.message}' for e in v.iter_errors(record)]


def controlla(records: Iterable[Mapping[str, Any]], v: Any | None = None) -> dict[str, Any]:
    """Every record against the schema, and the codes and anchors unique: counts and the first errors."""
    v = v or validatore()
    conti: Counter[str] = Counter()
    esempi: list[str] = []
    codici: Counter[str] = Counter()
    ancore: Counter[str] = Counter()
    for rec in records:
        conti['records'] += 1
        conti[f'corpus {rec.get("corpus")}'] += 1
        e = errori(rec, v)
        if e:
            conti['invalid'] += 1
            if len(esempi) < 20:
                esempi.append(f'{rec.get("ids", {}).get("code")}: {e[0]}')
        codici[rec.get('ids', {}).get('code')] += 1
        a = rec.get('ids', {}).get('anchor')
        if a:
            ancore[a] += 1
        for campo in ('document_import_id', 'family_import_id', 'case_id'):
            sezione = rec.get('cited' if campo == 'case_id' else 'citing') or {}
            if not sezione.get(campo):
                conti[f'without {campo}'] += 1
        if not (rec.get('citing') or {}).get('geography_iso'):
            conti['without geography_iso'] += 1
    conti['duplicate codes'] = sum(n - 1 for n in codici.values() if n > 1)
    conti['shared anchors'] = sum(n - 1 for n in ancore.values() if n > 1)
    return {'conti': dict(conti), 'esempi': esempi}


# ── command line ────────────────────────────────────────────────────────────────────────────────────────────────────


def _jsonl(percorso: Path) -> Iterator[dict[str, Any]]:
    with open(percorso, encoding='utf-8') as f:
        for riga in f:
            if riga.strip():
                yield json.loads(riga)


def converti(a: argparse.Namespace) -> Iterator[dict[str, Any]]:
    ids = Identificativi(Path(a.identificativi))
    pulizia = Pulizia()
    a._pulizia = pulizia
    if a.corpus == 'J':
        meta = MetadatiFonte(*(Path(x) for x in a.metadati))
        for r in _jsonl(Path(a.da)):
            yield da_j(r, ids, meta, pulizia)
    elif a.corpus == 'N':
        for r in _jsonl(Path(a.da)):
            yield da_n(r, ids, a.esecuzione, pulizia)
    else:
        indice = json.loads(Path(a.da).read_text(encoding='utf-8'))
        for x in indice.get('articoli', []):
            yield da_m(x, ids, indice.get('impronta'))


def aggiungi_parser(sub: Any) -> argparse.ArgumentParser:
    """Add ``traccia`` to the subcommands ``sub``: ``traccia converti`` and ``traccia valida``."""
    p = sub.add_parser('traccia', help='the common record of J, N and M (schema/trace-record.schema.json)')
    ps = p.add_subparsers(dest='cosa', required=True)
    pc = ps.add_parser('converti', help='a working file of one corpus as records (read only)')
    pc.add_argument('--corpus', required=True, choices=sorted(CORPORA))
    pc.add_argument('--da', required=True, help='J: citations .jsonl; N: S07-N rows .jsonl; M: indice.json')
    pc.add_argument('--identificativi', default=str(SCHEMA.parent.parent / 'identifiers'))
    pc.add_argument(
        '--metadati',
        action='append',
        default=[],
        help='J: metadata.csv or CCC-whole_data_dump-<date>.zip of the Climate Litigation Database, oldest first',
    )
    pc.add_argument('--esecuzione', default='S07-N-v3', help='N: the version code of the run')
    pc.add_argument('--uscita', required=True, help='the records, .jsonl')
    pc.add_argument('--valida', action='store_true', help='also check them against the schema')
    pv = ps.add_parser('valida', help='check a .jsonl of records against the schema')
    pv.add_argument('file', nargs='+')
    return p


def esegui(a: argparse.Namespace) -> int:
    if a.cosa == 'converti':
        uscita = Path(a.uscita)
        uscita.parent.mkdir(parents=True, exist_ok=True)
        records = list(converti(a))
        with open(uscita, 'w', encoding='utf-8') as f:
            for rec in records:
                f.write(json.dumps(rec, ensure_ascii=False) + '\n')
        print(len(records), 'records ->', uscita)
        if a._pulizia:
            print('control characters removed:', dict(a._pulizia))
        if not a.valida:
            return 0
        esito = controlla(records)
    else:
        esito = controlla(rec for p in a.file for rec in _jsonl(Path(p)))
    print(json.dumps(esito, ensure_ascii=False, indent=1))
    c = esito['conti']
    return 1 if c.get('invalid') or c.get('duplicate codes') else 0
